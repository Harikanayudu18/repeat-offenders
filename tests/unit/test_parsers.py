import io
import pytest
from backend.services.imports.parsers import ParseError, parse_upload

TXT = "YEAR 2021\n1. Explain the idea\n   of recursion. [5 marks]\n2. Define a stack (3)\n"


def test_txt_multiline_and_marks():
    r = parse_upload("p.txt", TXT.encode())
    assert [(x["year"], x["question_number"], x["marks"]) for x in r] == [("2021", "1", "5"), ("2021", "2", "3")]
    assert "recursion" in r[0]["question_text"] and "[5" not in r[0]["question_text"]


def test_xlsx_and_docx_roundtrip():
    import openpyxl, docx
    wb = openpyxl.Workbook(); wb.active.append(["Year", "Marks", "Question"]); wb.active.append([2020, 4, "Unicode ✓ α"])
    b = io.BytesIO(); wb.save(b)
    assert parse_upload("x.xlsx", b.getvalue())[0]["question_text"] == "Unicode ✓ α"
    d = docx.Document(); d.add_paragraph("YEAR 2019"); d.add_paragraph("1. Define entropy [6 marks]")
    b = io.BytesIO(); d.save(b)
    assert parse_upload("x.docx", b.getvalue())[0]["marks"] == "6"


def test_scanned_pdf_and_malformed_and_unsupported():
    from pypdf import PdfWriter
    w = PdfWriter(); w.add_blank_page(100, 100); b = io.BytesIO(); w.write(b)
    with pytest.raises(ParseError) as e: parse_upload("s.pdf", b.getvalue())
    assert e.value.code == "scanned_pdf"
    with pytest.raises(ParseError) as e: parse_upload("x.xlsx", b"not a zip")
    assert e.value.code == "malformed_file"
    with pytest.raises(ParseError) as e: parse_upload("../../evil.exe", b"x")
    assert e.value.status == 415
    with pytest.raises(ParseError) as e: parse_upload("e.csv", b"")
    assert e.value.code in ("empty_file",)
    with pytest.raises(ParseError) as e: parse_upload("m.csv", b"a,b\n1,2\n")
    assert e.value.code == "missing_columns"


def test_security_headers_and_host_check(client):
    r = client.get("/api/v1/health")
    assert r.headers["x-content-type-options"] == "nosniff" and r.headers["cache-control"] == "no-store"
    assert client.get("/api/v1/health", headers={"host": "evil.example"}).status_code == 400
