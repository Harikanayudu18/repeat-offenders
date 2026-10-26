"""Safe parsers: CSV, XLSX, DOCX, TXT, text-based PDF. Files are parsed in memory only; never written, executed or logged."""
import csv
import io
import re
import zipfile

MAX_ROWS, MAX_UNZIPPED, MAX_PDF_PAGES = 5000, 50_000_000, 300
REQUIRED = ("year", "marks", "question_text")
ALIASES = {"question": "question_text", "text": "question_text", "q": "question_number", "qno": "question_number",
           "number": "question_number", "mark": "marks", "q_no": "question_number"}


class ParseError(Exception):
    def __init__(self, status: int, code: str, message: str):
        super().__init__(message)
        self.status, self.code, self.message = status, code, message


def _hdr(h): k = str(h or "").strip().lower().replace(" ", "_"); return ALIASES.get(k, k)


def _tabular(rows):
    it = (r for r in rows if any(str(c or "").strip() for c in r))
    head = next(it, None)
    if head is None: raise ParseError(422, "empty_file", "The file contains no rows.")
    cols = [_hdr(h) for h in head]
    missing = [k for k in REQUIRED if k not in cols]
    if missing: raise ParseError(422, "missing_columns", f"Missing required columns: {missing}")
    out = []
    for r in it:
        out.append({c: (str(v).strip() if v is not None else "") for c, v in zip(cols, r)})
        if len(out) > MAX_ROWS: raise ParseError(422, "too_many_rows", f"More than {MAX_ROWS} rows.")
    return out


YEAR = re.compile(r"^\s*(?:year|paper|exam)?\s*[:\-]?\s*((?:19|20)\d{2})\s*$", re.I)
START = re.compile(r"^\s*(?:Q(?:uestion)?\s*)?(\d{1,3}[a-z]?)\s*[.):]\s+(.*)$", re.I)
MARKS = re.compile(r"[\[(]\s*(\d+(?:\.\d+)?)\s*(?:marks?|m)?\s*[\])]\s*$", re.I)


def parse_text(text: str, default_year: str = ""):
    """Format: optional 'YEAR 2021' lines; questions like '1. text ... [5 marks]' (may span lines)."""
    year, cur, out = default_year, None, []

    def flush():
        if cur:
            m = MARKS.search(cur["t"])
            out.append({"year": cur["y"], "question_number": cur["n"], "marks": m.group(1) if m else "",
                        "question_text": (MARKS.sub("", cur["t"]) if m else cur["t"]).strip(), "topic": ""})
    for ln in text.splitlines():
        if YEAR.match(ln): flush(); cur = None; year = YEAR.match(ln).group(1); continue
        m = START.match(ln)
        if m: flush(); cur = {"y": year, "n": m.group(1), "t": m.group(2).strip()}
        elif cur and ln.strip(): cur["t"] += " " + ln.strip()
    flush()
    if not out: raise ParseError(422, "no_questions", "No questions found. Use lines like '1. text [5 marks]' under a 'YEAR 2021' line.")
    return out[:MAX_ROWS]


def _zip_guard(raw):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED: raise ParseError(413, "archive_too_large", "File expands too much.")
    except zipfile.BadZipFile: raise ParseError(400, "malformed_file", "File is corrupt or not a valid Office document.")


def _decode(raw):
    try: return raw.decode("utf-8-sig")
    except UnicodeDecodeError: raise ParseError(400, "bad_encoding", "File must be UTF-8 text.")


def parse_upload(filename: str, raw: bytes):
    ext = (filename or "").lower().rsplit(".", 1)[-1] if "." in (filename or "") else ""
    year_hint = (re.search(r"(?:19|20)\d{2}", filename or "") or [""])[0]
    if ext == "csv":
        text = _decode(raw)
        try: d = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error: d = csv.excel
        return _tabular(list(csv.reader(io.StringIO(text), dialect=d)))
    if ext == "txt": return parse_text(_decode(raw), year_hint)
    if ext == "xlsx":
        _zip_guard(raw)
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            err = None
            for ws in wb.worksheets:  # first worksheet that has the required headers
                try: return _tabular(list(ws.iter_rows(values_only=True)))
                except ParseError as e: err = e
            raise err or ParseError(422, "empty_file", "No worksheets.")
        except ParseError: raise
        except Exception: raise ParseError(400, "malformed_file", "Could not read the workbook.")
    if ext == "docx":
        _zip_guard(raw)
        try:
            import docx
            d = docx.Document(io.BytesIO(raw))
            for t in d.tables:
                try: return _tabular([[c.text for c in r.cells] for r in t.rows])
                except ParseError: continue
            return parse_text("\n".join(p.text for p in d.paragraphs), year_hint)
        except ParseError: raise
        except Exception: raise ParseError(400, "malformed_file", "Could not read the Word document.")
    if ext == "pdf":
        try:
            from pypdf import PdfReader
            rd = PdfReader(io.BytesIO(raw))
            if len(rd.pages) > MAX_PDF_PAGES: raise ParseError(413, "too_many_pages", "PDF has too many pages.")
            pages = [(p.extract_text() or "") for p in rd.pages]
        except ParseError: raise
        except Exception: raise ParseError(400, "malformed_file", "Could not read the PDF.")
        if sum(len(p.strip()) for p in pages) < 30:
            raise ParseError(422, "scanned_pdf", "This PDF appears scanned or image-based (OCR is not supported). Paste the text into a .txt file or use a CSV.")
        return parse_text("\n".join(pages), year_hint)
    raise ParseError(415, "unsupported_type", "Supported formats: .csv, .xlsx, .docx, .txt, .pdf (text-based).")
